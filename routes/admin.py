from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, make_response
from flask_login import login_required, current_user
from models import db, User, Application, Report, Organization, Placement, Notification, Evaluation, Attendance
from datetime import datetime, date
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from openpyxl import Workbook
import io

admin_bp = Blueprint('admin', __name__)

def check_admin_access():
    if not current_user.is_authenticated or current_user.role != 'admin':
        flash('Access denied. Administrator account required.', 'error')
        return False
    return True

@admin_bp.route('/dashboard')
@login_required
def dashboard():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    # Get statistics
    total_students = User.query.filter_by(role='student').count()
    total_supervisors = User.query.filter_by(role='supervisor').count()
    total_organizations = Organization.query.count()
    pending_applications = Application.query.filter_by(status='pending').count()
    active_placements = Application.query.filter_by(status='in_progress').count()
    
    # Recent activities
    recent_applications = Application.query.order_by(Application.applied_at.desc()).limit(5).all()
    recent_reports = Report.query.order_by(Report.submitted_at.desc()).limit(5).all()
    
    stats = {
        'total_students': total_students,
        'total_supervisors': total_supervisors,
        'total_organizations': total_organizations,
        'pending_applications': pending_applications,
        'active_placements': active_placements,
        'total_applications': Application.query.count(),
        'total_reports': Report.query.count()
    }
    
    return render_template('admin/dashboard.html', 
                         stats=stats,
                         recent_applications=recent_applications,
                         recent_reports=recent_reports)

@admin_bp.route('/applications')
@login_required
def applications():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    status_filter = request.args.get('status', '')
    
    query = Application.query
    if status_filter:
        query = query.filter_by(status=status_filter)
    
    applications = query.order_by(Application.applied_at.desc()).all()
    return render_template('admin/applications.html', applications=applications, status_filter=status_filter)

@admin_bp.route('/application/<int:application_id>')
@login_required
def view_application(application_id):
    if not check_admin_access():
        return redirect(url_for('index'))
    
    application = Application.query.get_or_404(application_id)
    supervisors = User.query.filter_by(role='supervisor', is_active=True).all()
    
    return render_template('admin/view_application.html', application=application, supervisors=supervisors)

@admin_bp.route('/approve-application/<int:application_id>', methods=['POST'])
@login_required
def approve_application(application_id):
    if not check_admin_access():
        return redirect(url_for('index'))
    
    application = Application.query.get_or_404(application_id)
    action = request.form.get('action')
    supervisor_id = request.form.get('supervisor_id')
    start_date = request.form.get('start_date')
    end_date = request.form.get('end_date')
    admin_comments = request.form.get('admin_comments', '').strip()
    
    try:
        if action == 'approve':
            if not supervisor_id or not start_date or not end_date:
                flash('Please provide supervisor and placement dates', 'error')
                return redirect(url_for('admin.view_application', application_id=application_id))
            
            application.status = 'approved'
            application.supervisor_id = int(supervisor_id)
            application.start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
            application.end_date = datetime.strptime(end_date, '%Y-%m-%d').date()
            application.approved_at = datetime.utcnow()
            
        elif action == 'reject':
            application.status = 'rejected'
        
        elif action == 'start':
            application.status = 'in_progress'
        
        application.admin_comments = admin_comments
        db.session.commit()
        
        # Notify student
        notification_message = f'Your application has been {application.status}'
        if application.status == 'approved':
            supervisor = User.query.get(supervisor_id)
            notification_message += f' and assigned to supervisor {supervisor.full_name}'
        
        notification = Notification(
            recipient_id=application.student_id,
            title='Application Status Update',
            message=notification_message,
            type='success' if action in ['approve', 'start'] else 'warning'
        )
        db.session.add(notification)
        
        # Notify supervisor if assigned
        if action == 'approve' and supervisor_id:
            supervisor_notification = Notification(
                recipient_id=int(supervisor_id),
                title='New Student Assignment',
                message=f'You have been assigned to supervise {application.student.full_name}',
                type='info'
            )
            db.session.add(supervisor_notification)
        
        db.session.commit()
        
        flash(f'Application {application.status} successfully!', 'success')
        
    except Exception as e:
        db.session.rollback()
        flash('Failed to update application. Please try again.', 'error')
        print(f"Application update error: {e}")
    
    return redirect(url_for('admin.view_application', application_id=application_id))

@admin_bp.route('/students')
@login_required
def students():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    students = User.query.filter_by(role='student').order_by(User.created_at.desc()).all()
    return render_template('admin/students.html', students=students)

@admin_bp.route('/supervisors')
@login_required
def supervisors():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    supervisors = User.query.filter_by(role='supervisor').order_by(User.created_at.desc()).all()
    return render_template('admin/supervisors.html', supervisors=supervisors)

@admin_bp.route('/organizations')
@login_required
def organizations():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    organizations = Organization.query.order_by(Organization.created_at.desc()).all()
    return render_template('admin/organizations.html', organizations=organizations)

@admin_bp.route('/add-organization', methods=['GET', 'POST'])
@login_required
def add_organization():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        org_type = request.form.get('type', '').strip()
        location = request.form.get('location', '').strip()
        contact_person = request.form.get('contact_person', '').strip()
        email = request.form.get('email', '').strip()
        phone = request.form.get('phone', '').strip()
        description = request.form.get('description', '').strip()
        
        # Validation
        if not all([name, org_type, location]):
            flash('Please fill in all required fields', 'error')
            return render_template('admin/add_organization.html')
        
        try:
            organization = Organization(
                name=name,
                type=org_type,
                location=location,
                contact_person=contact_person,
                email=email,
                phone=phone,
                description=description
            )
            
            db.session.add(organization)
            db.session.commit()
            
            flash('Organization added successfully!', 'success')
            return redirect(url_for('admin.organizations'))
            
        except Exception as e:
            db.session.rollback()
            flash('Failed to add organization. Please try again.', 'error')
            print(f"Add organization error: {e}")
    
    return render_template('admin/add_organization.html')

@admin_bp.route('/placements')
@login_required
def placements():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    placements = Placement.query.order_by(Placement.created_at.desc()).all()
    return render_template('admin/placements.html', placements=placements)

@admin_bp.route('/add-placement', methods=['GET', 'POST'])
@login_required
def add_placement():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    if request.method == 'POST':
        title = request.form.get('title', '').strip()
        description = request.form.get('description', '').strip()
        organization_id = request.form.get('organization_id')
        location = request.form.get('location', '').strip()
        start_date = request.form.get('start_date')
        end_date = request.form.get('end_date')
        requirements = request.form.get('requirements', '').strip()
        available_slots = request.form.get('available_slots')
        
        # Validation
        if not all([title, organization_id, location, available_slots]):
            flash('Please fill in all required fields', 'error')
            return render_template('admin/add_placement.html', organizations=Organization.query.all())
        
        try:
            placement = Placement(
                title=title,
                description=description,
                organization_id=int(organization_id),
                location=location,
                requirements=requirements,
                available_slots=int(available_slots)
            )
            
            if start_date:
                placement.start_date = datetime.strptime(start_date, '%Y-%m-%d').date()
            if end_date:
                placement.end_date = datetime.strptime(end_date, '%Y-%m-%d').date()
            
            db.session.add(placement)
            db.session.commit()
            
            flash('Placement added successfully!', 'success')
            return redirect(url_for('admin.placements'))
            
        except Exception as e:
            db.session.rollback()
            flash('Failed to add placement. Please try again.', 'error')
            print(f"Add placement error: {e}")
    
    organizations = Organization.query.filter_by(is_active=True).all()
    return render_template('admin/add_placement.html', organizations=organizations)

@admin_bp.route('/reports')
@login_required
def reports():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    reports = Report.query.order_by(Report.submitted_at.desc()).all()
    return render_template('admin/reports.html', reports=reports)

@admin_bp.route('/generate-report')
@login_required
def generate_report():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    return render_template('admin/generate_report.html')

@admin_bp.route('/export-applications')
@login_required
def export_applications():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    format_type = request.args.get('format', 'excel')
    
    applications = Application.query.all()
    
    if format_type == 'pdf':
        return export_applications_pdf(applications)
    else:
        return export_applications_excel(applications)

def export_applications_excel(applications):
    output = io.BytesIO()
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Applications"
    
    # Headers
    headers = ['ID', 'Student Name', 'Email', 'Type', 'Status', 'Applied Date', 'Supervisor', 'Start Date', 'End Date']
    for col, header in enumerate(headers, 1):
        worksheet.cell(row=1, column=col, value=header)
    
    # Data
    for row, application in enumerate(applications, 2):
        worksheet.cell(row=row, column=1, value=application.id)
        worksheet.cell(row=row, column=2, value=application.student.full_name)
        worksheet.cell(row=row, column=3, value=application.student.email)
        worksheet.cell(row=row, column=4, value=application.application_type)
        worksheet.cell(row=row, column=5, value=application.status)
        worksheet.cell(row=row, column=6, value=application.applied_at.strftime('%Y-%m-%d'))
        worksheet.cell(row=row, column=7, value=application.supervisor.full_name if application.supervisor else 'Not Assigned')
        worksheet.cell(row=row, column=8, value=application.start_date.strftime('%Y-%m-%d') if application.start_date else '')
        worksheet.cell(row=row, column=9, value=application.end_date.strftime('%Y-%m-%d') if application.end_date else '')
    
    workbook.save(output)
    output.seek(0)
    
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    response.headers['Content-Disposition'] = 'attachment; filename=applications.xlsx'
    
    return response

def export_applications_pdf(applications):
    output = io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4)
    
    styles = getSampleStyleSheet()
    elements = []
    
    # Title
    title = Paragraph("Field Supervision Applications Report", styles['Title'])
    elements.append(title)
    elements.append(Spacer(1, 12))
    
    # Table data
    data = [['ID', 'Student', 'Type', 'Status', 'Applied Date']]
    
    for application in applications:
        data.append([
            str(application.id),
            application.student.full_name,
            application.application_type,
            application.status,
            application.applied_at.strftime('%Y-%m-%d')
        ])
    
    # Create table
    table = Table(data)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.grey),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 14),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black)
    ]))
    
    elements.append(table)
    doc.build(elements)
    output.seek(0)
    
    response = make_response(output.getvalue())
    response.headers['Content-Type'] = 'application/pdf'
    response.headers['Content-Disposition'] = 'attachment; filename=applications.pdf'
    
    return response

@admin_bp.route('/user/<int:user_id>/toggle-status', methods=['POST'])
@login_required
def toggle_user_status(user_id):
    if not check_admin_access():
        return redirect(url_for('index'))
    
    user = User.query.get_or_404(user_id)
    
    if user.role == 'admin':
        flash('Cannot deactivate admin accounts', 'error')
        return redirect(request.referrer or url_for('admin.dashboard'))
    
    try:
        user.is_active = not user.is_active
        db.session.commit()
        
        status = 'activated' if user.is_active else 'deactivated'
        flash(f'User {status} successfully!', 'success')
        
    except Exception as e:
        db.session.rollback()
        flash('Failed to update user status. Please try again.', 'error')
    
    return redirect(request.referrer or url_for('admin.dashboard'))

@admin_bp.route('/notifications')
@login_required
def notifications():
    if not check_admin_access():
        return redirect(url_for('index'))
    
    notifications = Notification.query.filter_by(recipient_id=current_user.id).order_by(Notification.created_at.desc()).all()
    
    # Mark notifications as read
    unread_notifications = [n for n in notifications if not n.is_read]
    for notification in unread_notifications:
        notification.is_read = True
    
    if unread_notifications:
        db.session.commit()
    
    return render_template('admin/notifications.html', notifications=notifications)